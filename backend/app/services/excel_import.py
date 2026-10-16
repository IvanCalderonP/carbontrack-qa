from dataclasses import dataclass
from io import BytesIO
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from pydantic import ValidationError

from app.core.errors import AppError
from app.schemas.contracts import ImportErrorRow, ProductInput

HEADERS = (
    "product_name",
    "cn_code",
    "country_of_origin",
    "quantity",
    "unit",
    "direct_emissions",
    "indirect_emissions",
    "carbon_price",
)
MAX_UPLOAD_BYTES = 2 * 1024 * 1024
MAX_DATA_ROWS = 1000


@dataclass
class ParsedWorkbook:
    total_rows: int
    products: list[tuple[int, ProductInput]]
    errors: list[ImportErrorRow]


def parse_products_workbook(content: bytes) -> ParsedWorkbook:
    """Validate a small .xlsx workbook without evaluating formulas or accepting macros."""
    if not content.startswith(b"PK"):
        raise AppError(400, "INVALID_WORKBOOK", "File is not a valid .xlsx workbook")
    workbook = None
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True, keep_links=False)
        sheet = workbook.active
        if sheet is None:
            raise ValueError("No active sheet")
        rows = sheet.iter_rows(values_only=True)
        header = next(rows, None)
        if header is None:
            raise AppError(422, "EMPTY_WORKBOOK", "Workbook has no header or data")
        actual = tuple(str(cell).strip() if cell is not None else "" for cell in header)
        if actual != HEADERS:
            raise AppError(422, "INVALID_HEADERS", "Headers must exactly match the documented columns")
        products: list[tuple[int, ProductInput]] = []
        errors: list[ImportErrorRow] = []
        seen: set[tuple[str, str, str]] = set()
        total_rows = 0
        for row_number, row in enumerate(rows, start=2):
            total_rows += 1
            if total_rows > MAX_DATA_ROWS:
                raise AppError(413, "TOO_MANY_ROWS", "Workbook exceeds 1000 data rows")
            if not any(cell is not None and str(cell).strip() for cell in row):
                errors.append(ImportErrorRow(row=row_number, field="row", message="Empty row"))
                continue
            data = dict(zip(HEADERS, row, strict=False))
            missing = [key for key in HEADERS if data.get(key) is None or str(data[key]).strip() == ""]
            if missing:
                errors.extend(
                    ImportErrorRow(row=row_number, field=key, message="Required cell is empty") for key in missing
                )
                continue
            data["name"] = data.pop("product_name")
            try:
                product = ProductInput.model_validate(data)
            except ValidationError as exc:
                errors.extend(
                    ImportErrorRow(row=row_number, field=str(issue["loc"][0]), message=issue["msg"])
                    for issue in exc.errors()
                )
                continue
            key = (product.name, product.cn_code, product.country_of_origin)
            if key in seen:
                errors.append(ImportErrorRow(row=row_number, field="product_name", message="Duplicate row"))
                continue
            seen.add(key)
            products.append((row_number, product))
        if total_rows == 0:
            raise AppError(422, "EMPTY_WORKBOOK", "Workbook has no data rows")
        return ParsedWorkbook(total_rows=total_rows, products=products, errors=errors)
    except AppError:
        raise
    except (
        OSError,
        ValueError,
        KeyError,
        IndexError,
        EOFError,
        TypeError,
        BadZipFile,
        ParseError,
        InvalidFileException,
    ):
        raise AppError(400, "INVALID_WORKBOOK", "File is not a valid .xlsx workbook") from None
    finally:
        if workbook is not None:
            workbook.close()
